import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

_HEADERS = [
    "Scenario ID",
    "Scenario Name",
    "Original Query",
    "Intent #",
    "Intent Description",
    "Domain",
    "Sub-Domain",
    "Data Source Type",
    "Business View (Structured)",
    "Unstructured Source",
    "Intent Types",
    "Data Fetch",
    "Reasoning",
    "Action",
]

_COL_WIDTHS = {
    1: 12, 2: 30, 3: 55, 4: 9,  5: 45,
    6: 14, 7: 20, 8: 18, 9: 38, 10: 35,
    11: 24, 12: 11, 13: 11, 14: 9,
}

_DOMAIN_ROW_COLORS = {
    "Sales":       "D6E4F7",
    "Finance":     "D6F5E3",
    "Operations":  "FFF3CD",
    "Logistics":   "F8D7DA",
    "Procurement": "E8D5F5",
    "Marketing":   "FCE4D6",
    "HR":          "D5E8D4",
}

_DEFAULT_ROW_COLOR = "F2F2F2"

_THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


class ReportGenerator:
    def generate(self, results: list[dict]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Intent Classification"

        self._write_headers(ws)
        row_idx = 2

        for result in results:
            s_id = result.get("scenario_id", "")
            s_name = result.get("scenario_name", "")
            query = result.get("original_query", "")
            intents = result.get("intents", [])

            for intent in intents:
                domain = intent.get("domain", "")
                fill_color = _DOMAIN_ROW_COLORS.get(domain, _DEFAULT_ROW_COLOR)
                row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

                values = [
                    s_id,
                    s_name,
                    query,
                    intent.get("intent_id", ""),
                    intent.get("description", ""),
                    domain,
                    intent.get("sub_domain", ""),
                    intent.get("data_source", ""),
                    intent.get("structured_view") or "",
                    intent.get("unstructured_source") or "",
                    ", ".join(intent.get("intent_types", [])),
                    "Yes" if intent.get("requires_data_fetch") else "No",
                    "Yes" if intent.get("requires_reasoning") else "No",
                    "Yes" if intent.get("requires_action") else "No",
                ]

                for col, val in enumerate(values, 1):
                    cell = ws.cell(row=row_idx, column=col, value=val)
                    cell.fill = row_fill
                    cell.border = _THIN_BORDER
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

                row_idx += 1

        for col, width in _COL_WIDTHS.items():
            ws.column_dimensions[get_column_letter(col)].width = width

        ws.freeze_panes = "A2"

        self._write_summary(wb, results)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    def _write_headers(self, ws) -> None:
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=10)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32

        for col, name in enumerate(_HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = _THIN_BORDER

    def _write_summary(self, wb: openpyxl.Workbook, results: list[dict]) -> None:
        ws = wb.create_sheet("Summary")
        bold = Font(bold=True)

        total_intents = sum(len(r.get("intents", [])) for r in results)
        domain_counts: dict[str, int] = {}
        source_counts: dict[str, int] = {}

        for r in results:
            for intent in r.get("intents", []):
                d = intent.get("domain", "Unknown")
                domain_counts[d] = domain_counts.get(d, 0) + 1
                s = intent.get("data_source", "Unknown")
                source_counts[s] = source_counts.get(s, 0) + 1

        ws.append(["Report Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws.append(["Model", "GEMMA 4 (26B)"])
        ws.append(["Agent", "ARIA v1.0"])
        ws.append([])
        ws.append(["Total Scenarios Processed", len(results)])
        ws.append(["Total Intents Decomposed", total_intents])
        ws.append(["Avg Intents per Scenario", round(total_intents / len(results), 2) if results else 0])
        ws.append([])
        ws.append(["--- Intents by Domain ---", ""])
        for domain, count in sorted(domain_counts.items(), key=lambda x: -x[1]):
            ws.append([domain, count])
        ws.append([])
        ws.append(["--- Intents by Data Source ---", ""])
        for src, count in sorted(source_counts.items(), key=lambda x: -x[1]):
            ws.append([src, count])

        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 20

        for row in ws.iter_rows():
            for cell in row:
                if cell.col_idx == 1 and cell.value and str(cell.value).startswith("---"):
                    cell.font = Font(bold=True, color="1F4E79")
                elif cell.col_idx == 1 and cell.row <= 8:
                    cell.font = bold
