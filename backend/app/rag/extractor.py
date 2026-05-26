"""
Text extraction from file bytes by extension.
Returns (text: str, page_count: int, extraction_method: str).
Raises ExtractorError on failure.
"""

import io
import os
import logging
from typing import Tuple

logger = logging.getLogger(__name__)


class ExtractorError(Exception):
    pass


def extract(file_bytes: bytes, filename: str) -> Tuple[str, int, str]:
    """
    Extract plain text from file bytes.

    Returns:
        (text, page_count, extraction_method)

    Raises:
        ExtractorError — unrecoverable extraction failure
    """
    ext = os.path.splitext(filename.lower())[1]
    try:
        if ext == '.pdf':
            return _extract_pdf(file_bytes)
        elif ext == '.docx':
            return _extract_docx(file_bytes)
        elif ext in ('.txt', '.md'):
            return _extract_text(file_bytes)
        elif ext == '.csv':
            return _extract_csv(file_bytes)
        elif ext == '.xlsx':
            return _extract_xlsx(file_bytes)
        else:
            raise ExtractorError(f"Unsupported extension: {ext}")
    except ExtractorError:
        raise
    except Exception as e:
        raise ExtractorError(f"Extraction failed for {filename}: {e}") from e


# ── PDF ───────────────────────────────────────────────────────────────────────
def _extract_pdf(file_bytes: bytes) -> Tuple[str, int, str]:
    """
    Extract text from PDF using Docling.
    Handles native-text, scanned (OCR), multi-column, and table-heavy PDFs.
    Falls back to pypdf if Docling is not installed.
    """
    try:
        import tempfile
        from docling.document_converter import DocumentConverter

        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        try:
            converter = DocumentConverter()
            result = converter.convert(tmp_path)
            doc = result.document
            text = doc.export_to_markdown()
            page_count = len(doc.pages) if hasattr(doc, 'pages') and doc.pages else 1
        finally:
            import os
            try:
                os.unlink(tmp_path)
            except OSError:
                pass

        if not text.strip():
            raise ExtractorError("No text extracted from PDF (empty document)")

        return text, page_count, 'docling'

    except ImportError:
        logger.warning("docling not installed — falling back to pypdf. Run: pip install docling")
        return _extract_pdf_pypdf(file_bytes)


def _extract_pdf_pypdf(file_bytes: bytes) -> Tuple[str, int, str]:
    """Fallback PDF extractor using pypdf (native text only, no OCR)."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ExtractorError("Neither docling nor pypdf is installed. Run: pip install docling")

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        txt = page.extract_text() or ''
        pages.append(txt.strip())

    text = '\n\n'.join(p for p in pages if p)
    if not text.strip():
        raise ExtractorError("No text extracted from PDF (possibly scanned — install docling for OCR support)")

    return text, len(reader.pages), 'pypdf'


# ── DOCX ──────────────────────────────────────────────────────────────────────
def _extract_docx(file_bytes: bytes) -> Tuple[str, int, str]:
    try:
        from docx import Document
    except ImportError:
        raise ExtractorError("python-docx not installed. Run: pip install python-docx")

    doc = Document(io.BytesIO(file_bytes))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = '\n\n'.join(paras)
    if not text.strip():
        raise ExtractorError("No text extracted from DOCX")

    return text, 1, 'python-docx'


# ── TXT / MD ──────────────────────────────────────────────────────────────────
def _extract_text(file_bytes: bytes) -> Tuple[str, int, str]:
    try:
        text = file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        text = file_bytes.decode('latin-1', errors='replace')

    if not text.strip():
        raise ExtractorError("File is empty")

    return text, 1, 'text'


# ── CSV ───────────────────────────────────────────────────────────────────────
def _extract_csv(file_bytes: bytes) -> Tuple[str, int, str]:
    import csv

    try:
        content = file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        content = file_bytes.decode('latin-1', errors='replace')

    reader = csv.reader(io.StringIO(content))
    rows = [', '.join(row) for row in reader if any(c.strip() for c in row)]
    text = '\n'.join(rows)

    if not text.strip():
        raise ExtractorError("CSV file is empty")

    return text, 1, 'csv'


# ── XLSX ──────────────────────────────────────────────────────────────────────
def _extract_xlsx(file_bytes: bytes) -> Tuple[str, int, str]:
    try:
        import openpyxl
    except ImportError:
        raise ExtractorError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else '' for c in row]
            if any(c.strip() for c in cells):
                rows.append(', '.join(cells))
        if rows:
            sheets.append(f"[Sheet: {sheet_name}]\n" + '\n'.join(rows))
    wb.close()

    text = '\n\n'.join(sheets)
    if not text.strip():
        raise ExtractorError("No content found in XLSX")

    return text, len(wb.sheetnames), 'openpyxl'
